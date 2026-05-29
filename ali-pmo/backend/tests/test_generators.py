from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from generators.output_generator import write_ms_project_xml, write_project_plan_xlsx, write_raid_log_xlsx
from generators.plan_generator import generate_project_plan


ANALYSIS = {
    "project_title": "Finance Transformation PMO",
    "scope": ["Implement consolidated reporting"],
    "objectives": ["Improve executive visibility"],
    "deliverables": ["Integrated project plan", "RAID log"],
    "milestones": ["Sponsor approval on 2026-06-15"],
    "tasks": [
        {
            "name": "Prepare project charter",
            "description": "Prepare project charter | Owner: PMO Lead",
            "owner": "PMO Lead",
            "status": "Not Started",
            "dependencies": [],
            "dates": ["2026-06-01"],
        },
        {
            "name": "Build reporting dashboard",
            "description": "Build reporting dashboard requires finance data feed",
            "owner": "TBD Owner",
            "status": "At Risk",
            "dependencies": ["requires finance data feed"],
            "dates": [],
        },
    ],
    "risks": [
        {
            "description": "Finance feed delay could impact pilot",
            "owner": "TBD Owner",
            "impact": "High",
            "probability": "TBD",
            "status": "Open",
            "escalation_required": True,
            "mitigation": "Confirm data feed owner",
        }
    ],
    "assumptions": ["Business owners available"],
    "dependencies": [
        {
            "description": "Pending finance data feed",
            "owner": "ERP Lead",
            "status": "Blocked",
            "escalation_required": True,
        }
    ],
    "acceptance_criteria": ["Approved by steering committee"],
    "missing_information": ["Constraints"],
}


def test_project_plan_generation() -> None:
    plan = generate_project_plan("file-1", ANALYSIS)

    assert plan["project_title"] == "Finance Transformation PMO"
    assert any(task["owner"] == "PMO Lead" for task in plan["tasks"])
    assert any(task["owner"] == "TBD Owner" for task in plan["tasks"])
    assert any(task["milestone"] for task in plan["tasks"])


def test_excel_generation(tmp_path: Path) -> None:
    plan = generate_project_plan("file-1", ANALYSIS)
    path = write_project_plan_xlsx(tmp_path / "project_plan.xlsx", plan)

    workbook = load_workbook(path)
    sheet = workbook["Project Plan"]

    assert sheet["A1"].value == "WBS ID"
    assert sheet.max_row >= 3


def test_raid_generation(tmp_path: Path) -> None:
    path = write_raid_log_xlsx(tmp_path / "raid_log.xlsx", ANALYSIS)

    workbook = load_workbook(path)
    sheet = workbook["RAID Log"]

    assert sheet["B2"].value == "Risk"
    assert sheet["J2"].value == "Yes"


def test_ms_project_xml_generation(tmp_path: Path) -> None:
    plan = generate_project_plan("file-1", ANALYSIS)
    path = write_ms_project_xml(tmp_path / "ms_project.xml", ANALYSIS, plan)

    tree = ET.parse(path)
    root = tree.getroot()

    assert root.tag.endswith("Project")
    assert root.find(".//{http://schemas.microsoft.com/project}Tasks") is not None
    assert root.find(".//{http://schemas.microsoft.com/project}Task/{http://schemas.microsoft.com/project}Name").text == "Finance Transformation PMO"
