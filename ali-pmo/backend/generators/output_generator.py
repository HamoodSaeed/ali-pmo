from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from storage.project_storage import output_dir, update_metadata


PROJECT_PLAN_COLUMNS = [
    ("wbs_id", "WBS ID"),
    ("phase", "Phase"),
    ("task_name", "Task Name"),
    ("description", "Description"),
    ("start_date", "Start Date"),
    ("end_date", "End Date"),
    ("duration", "Duration"),
    ("owner", "Owner"),
    ("status", "Status"),
    ("dependency_predecessor", "Dependency Predecessor"),
    ("milestone", "Milestone"),
    ("notes", "Notes"),
]

RAID_COLUMNS = [
    "ID",
    "Type",
    "Description",
    "Owner",
    "Impact",
    "Probability",
    "Status",
    "Due Date",
    "Mitigation / Action",
    "Escalation Required",
]


def generate_outputs(file_id: str, analysis: dict[str, Any], plan: dict[str, Any]) -> dict[str, Any]:
    folder = output_dir(file_id)
    folder.mkdir(parents=True, exist_ok=True)

    files = {
        "project_plan.xlsx": write_project_plan_xlsx(folder / "project_plan.xlsx", plan),
        "project_plan.csv": write_project_plan_csv(folder / "project_plan.csv", plan),
        "raid_log.xlsx": write_raid_log_xlsx(folder / "raid_log.xlsx", analysis),
        "executive_summary.md": write_executive_summary(folder / "executive_summary.md", analysis, plan),
        "weekly_status_update.md": write_weekly_status(folder / "weekly_status_update.md", analysis, plan),
        "ms_project.xml": write_ms_project_xml(folder / "ms_project.xml", analysis, plan),
    }

    outputs = [
        {
            "name": name,
            "path": str(path),
            "size_bytes": path.stat().st_size,
            "download_url": f"/api/download/{file_id}/{name}",
        }
        for name, path in files.items()
    ]
    update_metadata(file_id, {"outputs": outputs, "outputs_generated_at": datetime.now(timezone.utc).isoformat()})
    return {"file_id": file_id, "outputs": outputs}


def write_project_plan_xlsx(path: Path, plan: dict[str, Any]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Project Plan"
    sheet.append([label for _, label in PROJECT_PLAN_COLUMNS])
    style_header(sheet)

    for task in plan.get("tasks", []):
        sheet.append([task.get(key, "") for key, _ in PROJECT_PLAN_COLUMNS])
    autosize(sheet)
    workbook.save(path)
    return path


def write_project_plan_csv(path: Path, plan: dict[str, Any]) -> Path:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, label in PROJECT_PLAN_COLUMNS])
        for task in plan.get("tasks", []):
            writer.writerow([task.get(key, "") for key, _ in PROJECT_PLAN_COLUMNS])
    return path


def write_raid_log_xlsx(path: Path, analysis: dict[str, Any]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RAID Log"
    sheet.append(RAID_COLUMNS)
    style_header(sheet)

    row_id = 1
    for risk in analysis.get("risks", []):
        sheet.append(
            [
                f"RAID-{row_id:03d}",
                "Risk",
                risk.get("description", ""),
                risk.get("owner", "TBD Owner"),
                risk.get("impact", "Medium"),
                risk.get("probability", "TBD"),
                risk.get("status", "Open"),
                "TBD",
                risk.get("mitigation", "TBD mitigation / action"),
                "Yes" if risk.get("escalation_required") else "No",
            ]
        )
        row_id += 1

    for assumption in analysis.get("assumptions", []):
        sheet.append([f"RAID-{row_id:03d}", "Assumption", assumption, "TBD Owner", "Medium", "TBD", "Open", "TBD", "Validate assumption", "No"])
        row_id += 1

    for dependency in analysis.get("dependencies", []):
        sheet.append(
            [
                f"RAID-{row_id:03d}",
                "Dependency",
                dependency.get("description", ""),
                dependency.get("owner", "TBD Owner"),
                "High" if dependency.get("escalation_required") else "Medium",
                "TBD",
                dependency.get("status", "Open"),
                "TBD",
                "Track owner and target resolution date",
                "Yes" if dependency.get("escalation_required") else "No",
            ]
        )
        row_id += 1

    if row_id == 1:
        sheet.append(["RAID-001", "Assumption", "No RAID items extracted from the source document.", "TBD Owner", "TBD", "TBD", "Open", "TBD", "Review source document with PMO", "No"])

    autosize(sheet)
    workbook.save(path)
    return path


def write_executive_summary(path: Path, analysis: dict[str, Any], plan: dict[str, Any]) -> Path:
    lines = [
        "# Executive Summary",
        "",
        "## Project Overview",
        summary_sentence(analysis),
        "",
        "## Current Understanding",
        current_understanding(analysis),
        "",
        "## Key Deliverables",
        bullet_lines(analysis.get("deliverables"), "TBD deliverables"),
        "",
        "## Major Milestones",
        bullet_lines(analysis.get("milestones"), "TBD milestones"),
        "",
        "## Key Risks and Dependencies",
        bullet_lines([risk.get("description", "") for risk in analysis.get("risks", [])] + [dep.get("description", "") for dep in analysis.get("dependencies", [])], "No explicit risks or dependencies were extracted."),
        "",
        "## Missing Information",
        bullet_lines(analysis.get("missing_information"), "No major missing information detected by the rule-based extractor."),
        "",
        "## Recommended Next Actions",
        bullet_lines(recommended_actions(analysis, plan), "Review and validate generated project controls."),
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_weekly_status(path: Path, analysis: dict[str, Any], plan: dict[str, Any]) -> Path:
    today = datetime.now().date().isoformat()
    rows = phase_progress_rows(plan)
    lines = [
        f"Date: {today}",
        "Overall Status: Draft - Requires PM Review",
        f"Key Message: {analysis.get('project_title', 'Project')} has been structured into an initial PMO control plan; missing fields require validation.",
        "",
        "Progress:",
        "| Track | Progress | Current Status |",
        "| --- | --- | --- |",
        *rows,
        "",
        "Risks / Blockers:",
        bullet_lines([risk.get("description", "") for risk in analysis.get("risks", []) if risk.get("escalation_required")] or [dep.get("description", "") for dep in analysis.get("dependencies", []) if dep.get("escalation_required")], "No escalated risks or blockers detected."),
        "",
        "Next Actions:",
        bullet_lines(recommended_actions(analysis, plan), "Validate plan with project owner."),
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_ms_project_xml(path: Path, analysis: dict[str, Any], plan: dict[str, Any]) -> Path:
    namespace = "http://schemas.microsoft.com/project"
    ET.register_namespace("", namespace)
    project = ET.Element(f"{{{namespace}}}Project")
    ET.SubElement(project, f"{{{namespace}}}Name").text = plan.get("project_title", "Ali PMO Generated Project")
    ET.SubElement(project, f"{{{namespace}}}Title").text = plan.get("project_title", "Ali PMO Generated Project")
    ET.SubElement(project, f"{{{namespace}}}ScheduleFromStart").text = "1"
    ET.SubElement(project, f"{{{namespace}}}StartDate").text = first_task_date(plan, "start_date")
    ET.SubElement(project, f"{{{namespace}}}FinishDate").text = last_task_date(plan, "end_date")
    ET.SubElement(project, f"{{{namespace}}}CalendarUID").text = "1"

    calendars = ET.SubElement(project, f"{{{namespace}}}Calendars")
    calendar = ET.SubElement(calendars, f"{{{namespace}}}Calendar")
    ET.SubElement(calendar, f"{{{namespace}}}UID").text = "1"
    ET.SubElement(calendar, f"{{{namespace}}}Name").text = "Standard"
    ET.SubElement(calendar, f"{{{namespace}}}IsBaseCalendar").text = "1"

    tasks_element = ET.SubElement(project, f"{{{namespace}}}Tasks")
    root_task = ET.SubElement(tasks_element, f"{{{namespace}}}Task")
    ET.SubElement(root_task, f"{{{namespace}}}UID").text = "0"
    ET.SubElement(root_task, f"{{{namespace}}}ID").text = "0"
    ET.SubElement(root_task, f"{{{namespace}}}Name").text = plan.get("project_title", "Ali PMO Generated Project")
    ET.SubElement(root_task, f"{{{namespace}}}Type").text = "1"
    ET.SubElement(root_task, f"{{{namespace}}}IsNull").text = "0"
    ET.SubElement(root_task, f"{{{namespace}}}OutlineLevel").text = "0"

    wbs_to_uid: dict[str, int] = {}
    for index, task in enumerate(plan.get("tasks", []), start=1):
        uid = index
        wbs_to_uid[task["wbs_id"]] = uid
        task_element = ET.SubElement(tasks_element, f"{{{namespace}}}Task")
        ET.SubElement(task_element, f"{{{namespace}}}UID").text = str(uid)
        ET.SubElement(task_element, f"{{{namespace}}}ID").text = str(index)
        ET.SubElement(task_element, f"{{{namespace}}}Name").text = task.get("task_name", "TBD Task")
        ET.SubElement(task_element, f"{{{namespace}}}Type").text = "1"
        ET.SubElement(task_element, f"{{{namespace}}}IsNull").text = "0"
        ET.SubElement(task_element, f"{{{namespace}}}WBS").text = task.get("wbs_id", str(index))
        ET.SubElement(task_element, f"{{{namespace}}}OutlineNumber").text = task.get("wbs_id", str(index))
        ET.SubElement(task_element, f"{{{namespace}}}OutlineLevel").text = str(outline_level(task.get("wbs_id", "")))
        ET.SubElement(task_element, f"{{{namespace}}}Start").text = project_datetime(task.get("start_date"))
        ET.SubElement(task_element, f"{{{namespace}}}Finish").text = project_datetime(task.get("end_date"))
        ET.SubElement(task_element, f"{{{namespace}}}Duration").text = project_duration(task)
        ET.SubElement(task_element, f"{{{namespace}}}Milestone").text = "1" if task.get("milestone") else "0"
        ET.SubElement(task_element, f"{{{namespace}}}Notes").text = task.get("notes", "")
        predecessor = task.get("dependency_predecessor")
        if predecessor and predecessor in wbs_to_uid:
            link = ET.SubElement(task_element, f"{{{namespace}}}PredecessorLink")
            ET.SubElement(link, f"{{{namespace}}}PredecessorUID").text = str(wbs_to_uid[predecessor])
            ET.SubElement(link, f"{{{namespace}}}Type").text = "1"

    resources = ET.SubElement(project, f"{{{namespace}}}Resources")
    owners = sorted({task.get("owner") for task in plan.get("tasks", []) if task.get("owner") and task.get("owner") != "TBD Owner"})
    for index, owner in enumerate(owners, start=1):
        resource = ET.SubElement(resources, f"{{{namespace}}}Resource")
        ET.SubElement(resource, f"{{{namespace}}}UID").text = str(index)
        ET.SubElement(resource, f"{{{namespace}}}ID").text = str(index)
        ET.SubElement(resource, f"{{{namespace}}}Name").text = owner
        ET.SubElement(resource, f"{{{namespace}}}Type").text = "1"

    tree = ET.ElementTree(project)
    ET.indent(tree, space="  ")
    tree.write(path, encoding="utf-8", xml_declaration=True)
    return path


def style_header(sheet: Any) -> None:
    fill = PatternFill(fill_type="solid", fgColor="1F2937")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
    sheet.freeze_panes = "A2"


def autosize(sheet: Any) -> None:
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 48)


def summary_sentence(analysis: dict[str, Any]) -> str:
    title = analysis.get("project_title") or "TBD Project Title"
    scope = analysis.get("scope") or []
    if scope:
        return f"{title} is understood as a project focused on {scope[0]}."
    return f"{title} has been parsed into an initial project control structure. The source document does not provide a complete scope statement."


def current_understanding(analysis: dict[str, Any]) -> str:
    parts = []
    if analysis.get("background"):
        parts.append(analysis["background"])
    if analysis.get("objectives"):
        parts.append("Objectives include " + "; ".join(analysis["objectives"][:3]) + ".")
    if not parts:
        return "The document contains limited narrative context. Ali PMO extracted available control items and marked gaps as missing information."
    return " ".join(parts)


def bullet_lines(items: list[str] | None, fallback: str) -> str:
    cleaned = [str(item).strip() for item in (items or []) if str(item).strip()]
    if not cleaned:
        cleaned = [fallback]
    return "\n".join(f"- {item}" for item in cleaned)


def recommended_actions(analysis: dict[str, Any], plan: dict[str, Any]) -> list[str]:
    actions = [
        "Validate extracted scope, deliverables, milestones, and owners with the project sponsor.",
        "Confirm placeholder dates and replace all assumptions with approved baseline dates.",
        "Assign owners for all TBD Owner tasks and RAID items.",
    ]
    if analysis.get("risks") or analysis.get("dependencies"):
        actions.append("Review escalated risks and blocked dependencies in the RAID log.")
    if plan.get("missing_information"):
        actions.append("Close missing information gaps before baseline approval.")
    return actions


def phase_progress_rows(plan: dict[str, Any]) -> list[str]:
    phases: dict[str, list[dict[str, Any]]] = {}
    for task in plan.get("tasks", []):
        if "." in str(task.get("wbs_id", "")):
            phases.setdefault(task.get("phase", "Delivery"), []).append(task)
    rows = []
    for phase, tasks in phases.items():
        completed = sum(1 for task in tasks if task.get("status") == "Completed")
        progress = f"{completed}/{len(tasks)} tasks complete"
        status = "At Risk" if any(task.get("status") == "At Risk" for task in tasks) else "Draft"
        rows.append(f"| {phase} | {progress} | {status} |")
    return rows or ["| Project Controls | Draft plan generated | Requires PM review |"]


def first_task_date(plan: dict[str, Any], key: str) -> str:
    values = [task.get(key) for task in plan.get("tasks", []) if task.get(key)]
    return project_datetime(min(values) if values else datetime.now().date().isoformat())


def last_task_date(plan: dict[str, Any], key: str) -> str:
    values = [task.get(key) for task in plan.get("tasks", []) if task.get(key)]
    return project_datetime(max(values) if values else datetime.now().date().isoformat())


def project_datetime(value: str | None) -> str:
    return f"{value or datetime.now().date().isoformat()}T08:00:00"


def project_duration(task: dict[str, Any]) -> str:
    if task.get("milestone"):
        return "PT0H0M0S"
    hours = int(task.get("duration_days", 1)) * 8
    return f"PT{hours}H0M0S"


def outline_level(wbs: str) -> int:
    if not wbs:
        return 1
    if wbs.startswith("M"):
        return 1
    return wbs.count(".") + 1
