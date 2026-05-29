from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import fitz
from fastapi import HTTPException
from openpyxl import load_workbook

from app.config import PREVIEW_CHARACTER_LIMIT
from storage.project_storage import (
    extracted_text_path,
    get_metadata,
    load_extracted_text,
    save_extracted_text,
    save_parser_status,
    update_metadata,
    utc_now,
)


def parse_document(file_id: str, force: bool = False) -> dict[str, Any]:
    metadata = get_metadata(file_id)
    cached_text = load_extracted_text(file_id)
    if cached_text is not None and not force:
        return build_parse_response(cached_text, metadata, status="completed", cached=True)

    upload_path = Path(metadata["upload_path"])
    extension = metadata["extension"].lower()

    save_parser_status(file_id, "parsing", {"filename": metadata["original_filename"]})
    try:
        if extension == ".pdf":
            text, page_count = extract_pdf_text(upload_path)
        elif extension == ".txt":
            text, page_count = extract_txt_text(upload_path)
        elif extension == ".docx":
            text, page_count = extract_docx_text(upload_path)
        elif extension == ".csv":
            text, page_count = extract_csv_text(upload_path)
        elif extension == ".xlsx":
            text, page_count = extract_xlsx_text(upload_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported parser for {extension}")
    except Exception as exc:
        save_parser_status(file_id, "failed", {"error": str(exc)})
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=500, detail=f"Parsing failed: {exc}") from exc

    save_extracted_text(file_id, text)
    save_parser_status(file_id, "completed", {"page_count": page_count, "character_count": len(text)})
    update_metadata(
        file_id,
        {
            "page_count": page_count,
            "character_count": len(text),
            "extracted_text_path": str(extracted_text_path(file_id)),
            "parsed_at": utc_now(),
        },
    )
    return build_parse_response(text, {**metadata, "page_count": page_count}, status="completed", cached=False)


def build_parse_response(text: str, metadata: dict[str, Any], status: str, cached: bool) -> dict[str, Any]:
    return {
        "file_id": metadata["file_id"],
        "page_count": metadata.get("page_count", 1),
        "character_count": len(text),
        "extraction_status": status,
        "cached": cached,
        "preview_text": text[:PREVIEW_CHARACTER_LIMIT],
    }


def extract_pdf_text(path: Path) -> tuple[str, int]:
    pages: list[str] = []
    with fitz.open(path) as document:
        for index, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            pages.append(f"--- Page {index} ---\n{page_text}")
        return "\n\n".join(pages).strip(), document.page_count


def extract_txt_text(path: Path) -> tuple[str, int]:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding).strip(), 1
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore").strip(), 1


def extract_docx_text(path: Path) -> tuple[str, int]:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=501, detail="DOCX parsing requires python-docx") from exc

    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))
    return "\n".join(paragraphs + table_rows), 1


def extract_csv_text(path: Path) -> tuple[str, int]:
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(rows), 1


def extract_xlsx_text(path: Path) -> tuple[str, int]:
    workbook = load_workbook(path, data_only=True)
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        lines.append(f"Sheet: {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines), len(workbook.worksheets)
